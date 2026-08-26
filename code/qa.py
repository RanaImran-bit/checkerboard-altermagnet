"""Geometric QA with word wrap modelled.

Text boxes here have word_wrap on, so a long paragraph does not overflow
horizontally, it wraps and consumes height. The check that matters is therefore
total wrapped height against box height. Monospace code blocks do NOT wrap
meaningfully (a wrapped code line is a defect), so those are checked on width.
"""
import math
from pptx import Presentation
EMU = 914400.0

prs = Presentation("code_changes_checkerboard.pptx")
SW, SH = prs.slide_width/EMU, prs.slide_height/EMU
hard = 0
for n, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        x, y, w, h = sh.left/EMU, sh.top/EMU, sh.width/EMU, sh.height/EMU
        boxes.append((x, y, w, h))
        usable = w - 0.20
        mono = any((r.font.name or "") == "Courier New"
                   for p in sh.text_frame.paragraphs for r in p.runs)
        per_em = 0.60 if mono else 0.48
        total_lines, maxsz = 0, 0
        for p in sh.text_frame.paragraphs:
            txt = "".join(r.text for r in p.runs)
            sz = max([(r.font.size.pt if r.font.size else 18) for r in p.runs], default=18)
            maxsz = max(maxsz, sz)
            wid = len(txt) * (sz/72.0) * per_em
            if mono:
                if wid > usable:
                    hard += 1
                    print(f"  s{n}: CODE LINE TOO WIDE {wid:.2f}/{usable:.2f}in | {txt[:50]!r}")
                total_lines += 1
            else:
                total_lines += max(1, math.ceil(wid/usable)) if usable > 0 else 99
        need = total_lines * (maxsz/72.0) * 1.22 + 0.12
        if need > h + 0.02:
            hard += 1
            print(f"  s{n}: HEIGHT OVERFLOW {need:.2f}/{h:.2f}in ({total_lines} lines "
                  f"@ {maxsz:g}pt) | {sh.text_frame.text[:46]!r}")
    for i in range(len(boxes)):
        for j in range(i+1, len(boxes)):
            ax, ay, aw, ah = boxes[i]; bx, by, bw, bh = boxes[j]
            ox = min(ax+aw, bx+bw) - max(ax, bx)
            oy = min(ay+ah, by+bh) - max(ay, by)
            if ox > 0.02 and oy > 0.02:
                hard += 1
                print(f"  s{n}: TEXT OVERLAP {ox:.2f} x {oy:.2f} in")
print(f"\n{hard} issue(s)")
